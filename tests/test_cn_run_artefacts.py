"""CARD-12 §4 — run 目录里的「各模块产物」。

§4 把证据包的输入定义成「run 目录（``manifest.json`` + ``events.jsonl`` + 各模块
产物）」。前两件 recorder 一直在写，**第三件从来没人写**：

* pack 因此最多只能装 4/13 件，``missing_completeness`` 恒等于「差 8 件」；
* 唯一一个名叫 ``claims.json`` 的件是**手写的占位符**，断言"本演示 run 未产生
  claim 级产物" —— 而它旁边那条 run 明明注册了 5 条 claim。一个与邻居直接矛盾的
  占位符比缺件更糟：它读起来像一条结论。

本模块钉住补齐后的产物。三条底线：

1. 内容来自 run 自己的状态（``analysis.to_dict()`` / 图谱），不第二次序列化、
   不重算，读者拿 pack 和 UI 对照时看到的是同一份形状；
2. **不虚构** run 没产出的件 —— 文档级 run 没有 ``valuation.xlsx``，它就该缺，
   并由 :func:`src.cn.run_artefacts.declared_absences` 写明为什么；
3. 同一 run 导出两次必须逐字节一致（xlsx 的时间戳属于 run，不属于导出动作）。
"""

from __future__ import annotations

import json
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.cn import demo_service as service  # noqa: E402
from src.cn.evidence_pack import (  # noqa: E402
    CATEGORY_COMPLETENESS,
    CATEGORY_REQUIRED,
    PACK_FILES,
    PACK_FILE_NAMES,
)
from src.cn.run_artefacts import (  # noqa: E402
    RUN_ARTEFACT_NAMES,
    _METRIC_COLUMNS,
    build_artefacts,
    declared_absences,
    run_scope,
    write_run_artefacts,
)

RUN_ID = "run-for-artefact-tests"
#: run 起始时刻。xlsx 的时间戳必须跟着它走，而不是跟着「保存这一刻」。
STARTED_AT = "2026-09-27T00:00:00+08:00"
EXPECTED_STAMP = "2026-09-26T16:00:00Z"  # 同一时刻的朴素 UTC 表示


@pytest.fixture(scope="module")
def filing():
    """整份文档链只建一次 —— 解析是这批用例里唯一的重活。"""
    state = service._filing(service.DEFAULT_FILING_KEY)
    yield state
    service.reset()


@pytest.fixture(scope="module")
def run_dir(tmp_path_factory) -> Path:
    """一个带真 manifest 的 run 目录，让产物有 run 上下文可读。"""
    directory = tmp_path_factory.mktemp("run_artefacts")
    (directory / "manifest.json").write_text(
        json.dumps(
            {
                "run_id": RUN_ID,
                "started_at": STARTED_AT,
                "finished_at": "2026-09-27T00:00:09+08:00",
                "input_document_sha256": {"annual_06_equity_shares.pdf": "0" * 64},
                "parser_name": "financial_statement",
                "parser_version": "1.0.0",
                "execution_mode": "offline",
                "snapshot_id": None,
                "git_commit": "0" * 40,
                "random_seed": 0,
                "model": None,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return directory


@pytest.fixture(scope="module")
def artefacts(filing, run_dir) -> dict:
    return build_artefacts(filing, run_dir, RUN_ID)


# --------------------------------------------------------------------------- #
# A. 清单：写出来的每一件都得是 pack 声明过的件
# --------------------------------------------------------------------------- #
def test_artefact_names_are_all_declared_pack_files():
    """产物名 = pack 的文件名。多一个没人校验的名字等于悄悄多出一份影子产物。"""
    assert set(RUN_ARTEFACT_NAMES) <= set(PACK_FILE_NAMES)


def test_build_artefacts_produces_exactly_the_declared_names(filing, run_dir, artefacts):
    assert set(artefacts) == set(RUN_ARTEFACT_NAMES)


def test_no_valuation_workbook_is_invented(filing, run_dir, artefacts):
    """文档级 run 不该有估值表，而且**不许**补一个空表凑满 13 件。"""
    assert "valuation.xlsx" not in RUN_ARTEFACT_NAMES
    assert "valuation.xlsx" not in artefacts
    # 但它的缺席必须有出处，否则下一次读残包的人只能重新推一遍为什么。
    absences = declared_absences(filing)
    assert "valuation.xlsx" in absences
    assert absences["valuation.xlsx"].strip()


def test_scope_is_read_off_the_run_not_asserted_by_the_caller(filing):
    assert run_scope(filing) == "document"


def test_the_hand_written_claims_placeholder_is_gone(artefacts):
    """旧实现在 pack 里塞了一句"本演示 run 未产生 claim 级产物"。"""
    assert "未产生 claim 级产物" not in json.dumps(artefacts["claims.json"], ensure_ascii=False)


# --------------------------------------------------------------------------- #
# B. claims.json —— 与图谱逐条对得上
# --------------------------------------------------------------------------- #
def test_claims_payload_matches_the_graph(filing, artefacts):
    payload = artefacts["claims.json"]
    graph = filing.graph
    assert payload["run_id"] == RUN_ID
    assert payload["claim_count"] == len(graph.claims) > 0
    assert {row["claim_id"] for row in payload["claims"]} == set(graph.claims)
    assert sum(payload["by_type"].values()) == payload["claim_count"]
    assert sum(payload["by_validation_status"].values()) == payload["claim_count"]


def test_claim_rows_keep_the_links_the_graph_recorded(artefacts):
    for row in artefacts["claims.json"]["claims"]:
        for key in ("evidence_ids", "calculation_ids", "derived_from", "assumption_ids"):
            assert isinstance(row[key], list), f"{row['claim_id']}.{key} 应当是列表"
        assert {"structural_status", "propagated_status", "reasons"} <= set(row)


def test_blocked_by_is_the_trigger_string_not_its_characters(artefacts):
    """``claim_status`` 的第 3 个返回值是 ``"type:id:state"`` 字符串（或 None）。

    早期实现写了 ``list(blocked)``，于是 ``"claim:C1:blocked"`` 被拆成一串单字符。
    这里同时钉住类型与形态 —— 只钉 ``is not list`` 挡不住 ``list("...")`` 以外的写法。
    """
    for row in artefacts["claims.json"]["claims"]:
        blocked = row["blocked_by"]
        assert blocked is None or isinstance(blocked, str), (
            f"{row['claim_id']}.blocked_by 应当是字符串或 None，实际 {type(blocked).__name__}"
        )
        if isinstance(blocked, str):
            assert blocked.count(":") >= 1


# --------------------------------------------------------------------------- #
# C. evidence.json —— 账本原样搬，页码不许猜
# --------------------------------------------------------------------------- #
def test_evidence_payload_carries_every_ledger_record(filing, artefacts):
    payload = artefacts["evidence.json"]
    records = filing.ledger.records()
    assert payload["record_count"] == len(records) > 0
    assert {record["evidence_id"] for record in payload["records"]} == {
        record.evidence_id for record in records
    }
    assert payload["snapshot_id"] == filing.ledger.snapshot_id


def test_pages_we_do_not_have_are_listed_not_filled_in(filing, artefacts):
    """没有页码的证据进 ``page_unlocated``，不补一个猜出来的页码。"""
    payload = artefacts["evidence.json"]
    located = payload["page_locations"]
    unlocated = payload["page_unlocated"]

    assert set(located) == set(filing.graph.evidence_index)
    for evidence_id in unlocated:
        assert located[evidence_id]["page"] is None
        assert located[evidence_id]["jumpable"] is False
    for evidence_id, entry in located.items():
        if entry["page"] is not None:
            assert evidence_id not in unlocated
            assert entry["jumpable"] is True


# --------------------------------------------------------------------------- #
# D. calculations.json / validation.json
# --------------------------------------------------------------------------- #
def test_calculations_payload_matches_the_graph(filing, artefacts):
    payload = artefacts["calculations.json"]
    assert payload["calculation_count"] == len(filing.graph.calculations) > 0
    assert {row["calculation_id"] for row in payload["calculations"]} == set(
        filing.graph.calculations
    )
    for row in payload["calculations"]:
        assert row["formula"], f"{row['calculation_id']} 没有公式"
        for operand in row["inputs"]:
            # GraphOperand 的字段是 kind/ref/value/label —— 没有 unit。
            assert set(operand) == {"kind", "ref", "value", "label"}


def test_validation_blocked_list_is_derived_not_guessed(artefacts):
    claims = {row["claim_id"]: row for row in artefacts["claims.json"]["claims"]}
    blocked = artefacts["validation.json"]["blocked_claims"]
    assert {row["claim_id"] for row in blocked} <= set(claims)
    for row in blocked:
        source = claims[row["claim_id"]]
        assert source["propagated_status"] != "supported" or source["blocked_by"]
    assert sum(artefacts["validation.json"]["by_propagated_status"].values()) == len(claims)


# --------------------------------------------------------------------------- #
# E. report.{json,md} —— 同一组数字，两种长相
# --------------------------------------------------------------------------- #
def test_report_json_names_the_run_and_the_scope(filing, artefacts):
    report = artefacts["report.json"]
    assert report["run_id"] == RUN_ID
    assert report["scope"] == "document"
    assert report["document_id"] == filing.document_id
    assert report["case"]["key"] == filing.case.key
    # 每个产出都得能追回自己的 run，否则 pack 里就是第二个互不相干的命名空间。
    for name in RUN_ARTEFACT_NAMES:
        payload = artefacts[name]
        if isinstance(payload, dict):
            assert payload.get("run_id") == RUN_ID, f"{name} 没写 run_id"


def test_report_markdown_carries_every_metric_row(filing, artefacts):
    markdown = artefacts["report.md"]
    for row in filing.analysis.to_dict()["metrics"]:
        assert str(row["label"]) in markdown
        assert str(row["calculation_id"]) in markdown
    assert "结构上不产出的件" in markdown
    assert "valuation.xlsx" in markdown


def test_manifest_brief_reads_the_run_header_from_disk(artefacts):
    manifest = artefacts["report.json"]["manifest"]
    assert manifest["run_id"] == RUN_ID
    assert manifest["started_at"] == STARTED_AT
    assert manifest["input_document_count"] == 1


def test_manifest_brief_is_empty_when_there_is_no_manifest(filing, tmp_path):
    """没有 manifest 就返回空，不编一个 run 头出来。"""
    payload = build_artefacts(filing, tmp_path / "no_manifest", RUN_ID)
    assert payload["report.json"]["manifest"] == {}


# --------------------------------------------------------------------------- #
# F. financials.xlsx —— report.json 的表格孪生
# --------------------------------------------------------------------------- #
def test_financials_workbook_mirrors_report_json_metrics(artefacts):
    workbook = load_workbook(BytesIO(artefacts["financials.xlsx"]))
    sheet = workbook.worksheets[0]
    assert sheet.title == "分析指标"
    header = [cell.value for cell in sheet[1]]
    assert header == [label for _key, label in _METRIC_COLUMNS]

    metrics = artefacts["report.json"]["analysis"]["metrics"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    # 指标表只放指标行：口径/告警在第二张表，切数据区不需要「遇到空行就停」。
    assert len(rows) == len(metrics) > 0
    assert [row[0] for row in rows] == [row["metric"] for row in metrics]


def test_workbook_notes_sheet_keeps_scope_and_gaps(artefacts):
    workbook = load_workbook(BytesIO(artefacts["financials.xlsx"]))
    notes = [row for row in workbook.worksheets[1].iter_rows(values_only=True)]
    assert workbook.worksheets[1].title == "口径与告警"
    assert notes[0] == ("类别", "内容")
    kinds = {row[0] for row in notes[1:]}
    assert {"口径", "版本", "报告期", "数据完整"} <= kinds
    analysis = artefacts["report.json"]["analysis"]
    if analysis.get("missing"):
        assert "missing" in kinds
    if analysis.get("warnings"):
        assert "warning" in kinds


def test_workbook_timestamps_come_from_the_run_not_from_the_export(artefacts):
    """openpyxl 默认盖「保存这一刻」；这份表属于 run，所以时间戳必须是 run 的。

    ``properties.modified`` 设置了也没用（openpyxl 3.1.5 保存时会重写它），
    因此这条断言同时是那条 ``_freeze_core`` 改写逻辑的回归锁。标签带命名空间属性，
    所以只钉「属性之后的正文」。
    """
    core = ZipFile(BytesIO(artefacts["financials.xlsx"])).read("docProps/core.xml").decode("utf-8")
    assert f">{EXPECTED_STAMP}</dcterms:created>" in core
    assert f">{EXPECTED_STAMP}</dcterms:modified>" in core


def test_artefacts_are_byte_identical_when_rebuilt(filing, run_dir, artefacts):
    """同一 run 重建两次必须完全一致 —— 否则 pack 的 zip 每次都不同。"""
    rebuilt = build_artefacts(filing, run_dir, RUN_ID)
    assert set(rebuilt) == set(artefacts)
    for name in RUN_ARTEFACT_NAMES:
        assert rebuilt[name] == artefacts[name], f"{name} 重建后不一致"


# --------------------------------------------------------------------------- #
# G. 落盘 + 端到端：跑一次真 run，导出，只有 valuation.xlsx 该缺
# --------------------------------------------------------------------------- #
def test_write_run_artefacts_lands_every_file(filing, tmp_path):
    written = write_run_artefacts(tmp_path / "run", state=filing, run_id=RUN_ID)
    assert set(written) == set(RUN_ARTEFACT_NAMES)
    for name, path in written.items():
        assert path.is_file(), f"{name} 没落盘"
        assert path.stat().st_size > 0, f"{name} 是空文件"
    report = json.loads((tmp_path / "run" / "report.json").read_text(encoding="utf-8"))
    assert report["run_id"] == RUN_ID


def test_end_to_end_pack_is_valid_and_only_valuation_is_missing(monkeypatch, tmp_path):
    """录一次真 run → 导出证据包：valid、无 error，缺的只有口径性的一件。"""
    monkeypatch.setattr(service, "RUN_STORE_DIR", tmp_path / "runs")
    monkeypatch.setattr(service, "PACK_DIR", tmp_path / "packs")

    recorded = service.record_demo_run()
    assert recorded["status"] == service.STATUS_OK, recorded.get("reason")
    assert set(recorded["artefacts"]) == set(RUN_ARTEFACT_NAMES)

    view = service.export_pack_view(recorded["run_id"])
    assert view["status"] == service.STATUS_OK, view.get("reason")
    assert view["valid"] is True
    assert view["errors"] == []
    assert view["missing_required"] == []
    assert view["missing_completeness"] == ["valuation.xlsx"]
    assert view["run_scope"] == "document"
    assert "valuation.xlsx" in view["declared_absences"]

    pack_dir = Path(view["pack_dir"])
    claims = json.loads((pack_dir / "claims.json").read_text(encoding="utf-8"))
    assert claims["claim_count"] > 0
    # 必装件 + 完整性件全部进包，只有文档级 run 产不出的估值表不在。
    present = {entry["name"] for entry in view["entries"] if entry["present"]}
    required = {
        name
        for name, category, _d in PACK_FILES
        if category in (CATEGORY_REQUIRED, CATEGORY_COMPLETENESS)
    }
    assert required - present == {"valuation.xlsx"}


def test_a_run_without_artefacts_is_reported_as_a_real_gap(monkeypatch, tmp_path):
    """run 目录里只有 manifest+events 时，包里缺的是**真缺口**，不许被说成口径。

    这是上面那条的反面：``declared_absences`` 来自 run 产物本身，产物不在就没有
    归因，于是 ``report.json`` 读不到、声明为空 —— 校验层照样只报「缺件」，但这次
    没有任何东西替它解释，只能当成该补的活。
    """
    from src.cn.runs.models import RunConfig
    from src.cn.runs.recorder import RunRecorder

    store = tmp_path / "runs"
    recorder = RunRecorder(store)
    run_id = recorder.start(
        RunConfig(
            input_documents=[],
            parser_name="bare",
            parser_version="0",
            execution_mode="offline",
        )
    )
    recorder.emit("claim", "claim_registered", {}, {}, "ok")
    recorder.finish("ok")

    monkeypatch.setattr(service, "RUN_STORE_DIR", store)
    monkeypatch.setattr(service, "PACK_DIR", tmp_path / "packs")

    view = service.export_pack_view(run_id)
    assert view["status"] == service.STATUS_OK, view.get("reason")
    assert view["declared_absences"] == {}
    assert view["run_scope"] is None
    assert "claims.json" in view["missing_completeness"]
    assert view["complete"] is False
