from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("docs/FinTrace-CN_研报纠错核查智能体_竞赛提交计划书.docx")

# Formal academic-paper palette: graphite anchors the tables, with a warm
# paper-gray alternating tint. It reads as a human-made competition plan,
# rather than a product or technology deck.
INK_BLUE = "3F464C"
MIST_BLUE = "F2F0EC"
CODE_MIST = "F0EEEA"
FOOTER_GRAY = "77736D"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def clear_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_text(doc, text, style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), 10.5, True)
        set_font(p.add_run(text[len(bold_lead):]), 10.5)
    else:
        set_font(p.add_run(text), 10.5)
    return p


def add_bullet(doc, text, level=0, compact=False):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(0 if compact else 3)
    p.paragraph_format.line_spacing = 1.0 if compact else 1.25
    set_font(p.add_run(text), 9.5 if compact else 10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), 10.5)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, INK_BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(value), 9.5, True, "FFFFFF")
        if widths:
            cell.width = Cm(widths[i])
    for r_idx, row_values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for i, value in enumerate(row_values):
            cell = cells[i]
            if r_idx % 2:
                set_cell_shading(cell, MIST_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            set_font(p.add_run(str(value)), 8.8)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_font(run, 15 if level == 1 else 12, True, "000000")
    return p


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_MIST)
    set_cell_margins(cell, 160, 180, 160, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("FinTrace-CN 竞赛实施计划")
    set_font(run, 8, False, FOOTER_GRAY)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    add_footer(section)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title_style_ppr = styles["Title"]._element.find(qn("w:pPr"))
    if title_style_ppr is not None:
        title_style_border = title_style_ppr.find(qn("w:pBdr"))
        if title_style_border is not None:
            title_style_ppr.remove(title_style_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    clear_paragraph_border(title)
    set_font(title.add_run("FinTrace CN 可追溯金融研究报告智能核查系统"), 18, True, "000000")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    set_font(subtitle.add_run("北京市大学生金融人工智能竞赛  初赛项目计划书"), 11, False, "555555")

    add_table(doc, ["项目字段", "内容"], [
        ["项目名称", "FinTrace CN 可追溯金融研究报告智能核查系统"],
        ["队伍名称", "〈待填〉"],
        ["学校 / 学院", "〈待填〉"],
        ["团队成员", "〈待填，3 至 5 人〉"],
        ["指导老师", "〈待填，晋级决赛后配备〉"],
        ["文档版本", "v2.2 · 2026-10-15"],
    ], [4.0, 13.0])

    add_heading(doc, "一 项目摘要")
    add_text(doc, "FinTrace CN 面向 A 股研究报告中最容易造成投资判断偏差的环节：财务数据、指标口径、估值倍数与引用依据在多份材料之间发生错配。系统以研究报告草稿和原始财务材料为输入，由大语言模型识别待核查主张并组织工具调用，再由本地确定性程序完成数据比对、期间与单位换算、公式复算、证据定位和修改建议生成。")
    add_text(doc, "项目不把大模型生成的文字直接当作金融事实。每项确认问题都必须包含研报原句、来源材料、页码或证据编号、核验规则或计算过程，以及可直接执行的修改建议。材料缺失、口径无法对齐或模型提取不充分时，系统明确输出“待人工复核”或“无法核验”。", bold_lead="项目不把")
    add_text(doc, "本项目主选赛题为“研究报告纠错核查”，并以“上市公司财务报告分析”提供事实底座，以“自动化估值建模”提供可验证的估值复算能力。最终交付物是一套可本地运行的智能核查原型、一组可复现样例、结构化核查报告与完整审计记录。")

    add_heading(doc, "二 业务需求与问题定义")
    add_heading(doc, "1 真实投研工作中的核查负担", 2)
    add_text(doc, "研究报告通常同时引用年报、季报、临时公告、行业材料和估值假设。人工复核需要在不同文件、不同期间、不同单位和不同统计口径之间反复跳转。常见失误并非复杂建模错误，而是“亿元和万元混用”“年报与单季混用”“归母与扣非口径混用”“同比公式写错”“估值分子分母不一致”以及“引用无法回溯”。这类问题会直接削弱研究结论的可信度。")
    add_heading(doc, "2 现有方法的不足", 2)
    add_table(doc, ["方法", "主要不足", "本项目的改进"], [
        ["人工逐页核查", "专业性强但耗时高，难保留完整复核轨迹", "将重复比对、换算与定位交给程序，保留人工判断"],
        ["通用 OCR 或文档问答", "可读文本但难判断期间、口径与计算关系", "建立金融事实规范化与确定性校验链"],
        ["纯大模型审稿", "可能产生幻觉，结果难解释、难复现", "模型只做主张提取和任务编排，不直接裁决事实"],
    ], [3.4, 6.4, 7.2])
    add_heading(doc, "3 项目目标", 2)
    for item in [
        "在受控数据环境中，自动从研究报告草稿中识别可验证的财务、估值和引用主张。",
        "对数值、单位、期间、口径、计算、估值与引用等高频问题形成带证据的核查结论。",
        "将模型推理、工具调用、数据访问、计算过程和最终结果全部记录，支持复核与复现。",
        "输出面向研究员的修改建议，而非仅输出“对”或“错”的聊天式回答。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "4 与竞赛评审要求的对应关系", 2)
    add_table(doc, ["竞赛关注点", "本项目的可见交付", "评委核验方式"], [
        ["任务完成度", "从材料登记到结构化核查报告的完整流程", "现场输入样例，查看任务轨迹与最终报告"],
        ["数据与计算准确性", "财务事实对象、单位期间约束、程序化公式与人工金标", "抽取任一错误，回看输入、公式和结果"],
        ["分析逻辑与专业性", "区分事实、计算错误、证据不足和需人工判断的观点", "查看错误类别、适用边界与修改建议"],
        ["可追溯与可复现", "文件哈希、页码、Evidence ID、规则与运行日志", "替换草稿或重复运行，核查结果链条"],
        ["创新与应用价值", "大模型负责理解编排，确定性规则负责事实裁决", "对比纯 OCR 或纯聊天式审稿的输出"],
    ], [3.2, 8.2, 5.6])

    add_heading(doc, "三 项目定位与范围")
    add_heading(doc, "1 项目定位", 2)
    add_text(doc, "项目名称：FinTrace-CN 面向 A 股研究报告的可追溯智能核查 Agent。")
    add_text(doc, "项目定义：系统读取原始财务材料和研究报告草稿，由大语言模型驱动的 Agent 识别核查任务并调用受控工具，完成事实提取、口径统一、财务与估值复算、引用定位和结果生成。每一条确认错误均可回到原文、数据、公式和执行记录。")
    add_heading(doc, "2 选题范围", 2)
    add_table(doc, ["层级", "竞赛选题", "在系统中的作用", "处理原则"], [
        ["主线", "5 研究报告纠错核查", "最终业务闭环与展示重点", "必须完成并稳定演示"],
        ["支撑", "2 上市公司财务报告分析", "提供收入、利润、现金流、期间和口径事实", "必须保留为核查依据"],
        ["加分", "4 自动化估值建模", "复算 PE PB PS 与可验证目标价输入", "仅做核查 不做完整 DCF"],
        ["暂缓", "买方投资备忘录、研究报告质量评估、产业链影响分析", "可作为后续扩展方向", "不进入初赛主流程，避免功能分散"],
    ], [2.0, 3.3, 7.0, 4.0])
    add_heading(doc, "3 明确不纳入初赛主线的内容", 2)
    for item in ["每日复盘与行情产品", "买方投资备忘录", "产业链影响分析", "多智能体和 MCP 服务化", "完整 DCF 流水线", "投资逻辑脆弱性和时间重验证", "对抗评测的产品化展示"]:
        add_bullet(doc, item)

    add_heading(doc, "四 核心创新与技术路线")
    add_heading(doc, "1 核心创新", 2)
    add_table(doc, ["创新点", "具体设计", "解决的竞赛关键问题"], [
        ["受控智能体", "模型负责识别主张、拆分任务与选择工具；确定性模块负责事实、计算和放行", "避免以自然语言回答替代金融核验"],
        ["金融事实规范化", "每个数值同时记录期间、单位、币种、口径、来源与证据定位", "避免跨材料直接比较导致的伪错误"],
        ["可计算的证据链", "错误结论绑定原句、原始材料位置、输入数据、公式和结果", "让修改建议可验证、可追溯"],
        ["失败状态可见", "材料不足、解析失败、口径冲突或模型超时时输出 blocked、partial 或人工复核", "不以“未发现错误”冒充“核验通过”"],
    ], [3.5, 7.7, 5.8])
    add_heading(doc, "2 端到端流程", 2)
    add_code(doc, [
        "原始年报或公告  +  研究报告草稿  +  研究截止日",
        "                ↓",
        "文件登记与哈希  →  文档解析  →  财务事实规范化",
        "                ↓",
        "主张提取  →  Agent 任务规划  →  受控工具调用",
        "                ↓",
        "确定性核验与复算  →  证据定位  →  修改建议  →  审计记录",
    ])
    add_heading(doc, "3 Agent 的职责边界", 2)
    add_text(doc, "Agent 的作用不是代替分析师下结论，而是把自然语言草稿转化为一组可执行、可验证的核查任务。它必须在预定义工具集合和步数上限内运行，所有工具返回统一状态：ok、partial、blocked 或 error。", bold_lead="Agent 的作用不是")
    add_heading(doc, "Agent 做什么", 2)
    for item in ["识别研报中收入、利润、同比、估值、引用、因果陈述等待核查主张。", "将主张拆成可验证任务，并在允许的工具集合中选择下一步。", "材料缺失、期间不明确或证据不足时，输出无法核验或待人工复核。", "将已经验证的结果组织为结构化核查报告和简明修改建议。"]:
        add_bullet(doc, item)
    add_heading(doc, "Agent 不做什么", 2)
    for item in ["不自行创造财务数字、来源、页码或证据。", "不自行计算同比、环比、PE、PB、PS 或单位换算。", "不自行决定事实是否成立，或绕过 Validator 输出结论。", "不将预测、分析师观点和可验证历史事实混为一谈。"]:
        add_bullet(doc, item)
    add_heading(doc, "执行控制", 2)
    add_text(doc, "采用单 Agent、受控工具调用和有限步数执行。建议每个任务最多 8 至 12 次工具调用，默认在本地材料和版本化快照中运行，不依赖互联网。模型负责任务理解与编排，确定性代码负责事实、计算、证据和最终放行。")

    add_heading(doc, "五 系统架构与核心能力")
    add_table(doc, ["能力层", "主要功能", "输出或控制点"], [
        ["材料登记", "登记文件哈希、来源、材料类型与研究截止日", "受控数据范围与可复现输入快照"],
        ["文档解析", "解析 PDF、复杂表格、扫描件与研报草稿", "文本片段、表格位置与解析状态"],
        ["金融事实处理", "抽取收入、利润、现金流、股本等事实并统一单位、期间和口径", "可比较的标准化金融事实"],
        ["主张理解与任务规划", "从研报原句中识别待核查主张，并组织受控工具调用", "结构化核查任务与调用轨迹"],
        ["确定性核验", "完成数值比对、公式复算、估值核验和引用检查", "错误类型、正确值、规则与严重等级"],
        ["证据与报告", "定位原文页码和证据编号，生成修改建议及结构化报告", "可回跳证据、审计记录和导出文件"],
    ], [3.3, 8.1, 7.6])
    add_heading(doc, "1 每条核查结果的必要字段", 2)
    add_table(doc, ["字段", "说明"], [
        ["finding_id 与 check_type", "唯一编号和错误类型，例如 unit_error 或 period_error"],
        ["draft_claim 与 draft_location", "研报原句及其页码或段落位置"],
        ["status 与 severity", "确认错误 警告 无法核验 观点预测 核验通过及严重等级"],
        ["expected_value", "正确数值 单位 期间和口径"],
        ["evidence", "原始材料 document_id 页码和 Evidence ID"],
        ["calculation 与 suggestion", "公式、输入、程序化复算结果与可直接执行的修改建议"],
    ], [5.4, 11.6])
    add_heading(doc, "2 首批稳定覆盖的错误类型", 2)
    add_text(doc, "初赛主展示优先覆盖八类：数值错误、单位错误、期间错误、统计口径错误、同比或环比计算错误、估值倍数复算错误、缺失或错误引用、错误页码或来源定位失败。")
    add_text(doc, "后续加分类型包括：过期引用、重述后仍使用旧值、因果断言证据不足、预测值与历史事实混淆、金融机构估值方法不适用。")

    add_heading(doc, "六 产品界面与交付形态")
    add_table(doc, ["页面", "评委可见内容", "主要目标"], [
        ["案例与材料", "原始材料 草稿 文件哈希 数据版本 研究截止日", "明确受控数据来源"],
        ["智能核查", "任务计划 工具调用 进度 成功或警告状态", "证明系统是 Agent 而非普通 OCR"],
        ["核查结果", "错误类型 原句 正确值 原因 修改建议", "形成业务交付物"],
        ["证据与复算", "PDF 页码 Evidence ID 规范化过程 公式", "证明结论可回溯"],
        ["审计轨迹与导出", "文件访问 工具记录 计算记录 Prompt 版本 输出哈希", "证明过程可复现"],
    ], [3.2, 10.5, 3.3])
    add_text(doc, "竞赛版应隐藏研究总览、每日复盘、市场行情、独立同行估值、任务中心、评测与消融及实验 Demo 页面。它们可以保留代码兼容性，但不应分散评委注意力。")

    add_heading(doc, "七 技术可行性与工程实现")
    add_text(doc, "原型采用“模型理解 + 受控工具 + 确定性规则 + 证据留痕”的分层实现。该设计避免把大模型作为唯一事实来源，同时保留模型处理非结构化语言和多步骤任务的优势。核心处理在本地或半封闭环境运行，第三方模型仅在被允许的范围内用于主张抽取与任务编排。")
    add_table(doc, ["工程层", "实现重点", "竞赛价值"], [
        ["数据层", "文档登记、文件哈希、解析结果缓存和版本化材料快照", "保证输入可核验，避免同名文件或版本漂移"],
        ["金融规则层", "单位、期间、口径、同比环比、估值倍数和引用核验规则", "保证结论可计算、可解释，不依赖模型猜测"],
        ["Agent 编排层", "受限工具清单、步数上限、状态机和失败处理", "体现智能体的任务组织能力与执行审慎性"],
        ["证据与审计层", "原句、页码、证据编号、公式输入输出、Prompt 与运行日志", "满足竞赛对可追溯、可复现的要求"],
        ["交互展示层", "材料、智能核查、核查结果、证据复算、审计导出五个视图", "让评委在短时间内看见完整闭环"],
    ], [3.3, 7.8, 7.9])
    add_heading(doc, "1 技术选型原则", 2)
    for item in [
        "模型可替换：通过统一接口接入国产或兼容模型，提交材料中明确模型名称、版本、调用方式和使用范围。",
        "规则可验证：所有会影响数值、期间、单位、口径和估值结论的处理均由确定性代码执行。",
        "依赖可复现：提供依赖文件、环境说明、示例数据与一键运行命令；不打包密钥、商业软件源代码或受版权限制材料。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "八 评测设计与演示数据")
    add_text(doc, "准备一套完全离线的固定演示材料，并为每个案例人工维护金标预期结果。合成植入错误集只能说明合成测试表现，不得表述为真实研报准确率。")
    add_table(doc, ["金标案例", "重点验证"], [
        ["正确研报", "零误报 不为展示效果而乱报错"],
        ["单位错误", "元 万元 亿元的换算与正确修改建议"],
        ["期间错误", "年报 半年报 单季 TTM 混用识别"],
        ["口径错误", "合并 母公司 归母 扣非等统计口径识别"],
        ["计算错误", "同比 环比或 PE PB 等程序化复算"],
        ["引用错误或数据过期", "来源 页码 重述和无法核验状态"],
    ], [5.3, 11.7])

    add_heading(doc, "1 评测原则与指标", 2)
    add_text(doc, "评测分为“系统能力评测”和“现场演示验证”两层。系统能力评测使用人工标注的固定案例，评估模型是否完整提取待核查主张，以及确定性校验器是否正确识别错误、控制误报并给出可追溯证据。现场演示验证强调一次真实流程的稳定运行、证据可回跳和失败状态可见。")
    add_table(doc, ["维度", "指标", "核验方式"], [
        ["主张提取", "可核查主张覆盖率、无效主张率", "与人工金标逐项比对，单列无法解析内容"],
        ["错误识别", "错误召回率、错误精确率、误报数", "按错误类型统计，不将警告混入确认错误"],
        ["数值与计算", "数值一致率、单位和期间换算正确率、公式复算正确率", "由程序和人工抽检双重确认"],
        ["证据可追溯", "证据定位成功率、页码或 Evidence ID 完整率", "随机抽取结果回跳至原始材料"],
        ["运行稳定性", "离线复跑成功率、失败状态显式率、日志完整率", "在同一版本数据和依赖下重复运行"],
    ], [3.0, 6.6, 7.4])
    add_heading(doc, "2 金标构建与边界控制", 2)
    add_text(doc, "每个案例由团队成员先对照原始财务材料建立人工金标，再由另一名成员复核错误类别、正确值、期间、单位和证据位置。对于预测、主观判断、产业观点等不可由原始材料直接验证的表述，金标不强行标记为“错误”，而是标记为“观点”“预测”或“需补充依据”。这保证评测只评价系统承诺解决的问题。")

    add_heading(doc, "九 已完成的受控运行案例")
    add_text(doc, "已完成一条受控端到端核查：以本地 2024 年年报作为参考材料，以自建含错研究草稿为待核查对象；DeepSeek 模型仅承担逐句 Claim 提取，数值比对、单位与期间换算、引用核验、错误判定和证据定位均由本地确定性模块完成。")
    add_table(doc, ["运行维度", "结果"], [
        ["模型调用", "DeepSeek V4.1 Flash，提取阶段按 6 句一批运行"],
        ["提取覆盖", "10 条 Claim，未提取 0；7 句为标题或非数值叙述"],
        ["核查输出", "8 项错误、9 项风险提示；保留原句、错误类型、建议与证据索引"],
        ["代表性问题", "收入数值、利润单位、同比、口径、估值、引用、研发费用期间"],
        ["可复现性", "JSON 记录模型与覆盖状态"],
    ], [4.3, 12.7])
    add_text(doc, "该案例用于证明系统的端到端可运行性与结果可追溯性，不将单个受控案例表述为通用模型准确率。")
    add_heading(doc, "1 案例的可复现条件", 2)
    add_table(doc, ["要素", "处理方式"], [
        ["输入材料", "以文件哈希、材料类型、研究截止日登记；现场使用可授权的样例材料"],
        ["模型角色", "模型名称、调用范围和批处理参数写入运行记录；模型不直接产出最终错误结论"],
        ["确定性规则", "单位、期间、口径、计算和引用核验由本地代码执行，并保留规则版本"],
        ["输出结果", "JSON 同时记录结果状态、主张覆盖、错误分类、证据索引和运行时间"],
    ], [3.4, 13.6])
    add_text(doc, "现场演示将提供一份“正确草稿”和一份“含错草稿”。前者用于展示误报控制，后者用于展示从主张提取、工具调用到错误回溯的完整闭环；两者均不依赖实时行情或外部网络。")

    add_heading(doc, "十 实施计划与里程碑")
    add_table(doc, ["阶段", "重点工作", "阶段交付物"], [
        ["范围冻结", "明确核查对象、错误类型、数据边界与不纳入范围", "范围说明、材料准入清单与人工金标规则"],
        ["主流程固化", "固定从材料登记、主张提取、工具调用到核查报告的端到端流程", "一套可重复运行案例、工具轨迹与结构化核查报告"],
        ["界面与证据完善", "完成原句、页码、计算过程和证据索引的联动展示", "五页竞赛工作台与可回跳证据链"],
        ["评测与演示", "执行离线测试、核验金标结果、准备正确和含错两类草稿", "测试记录、演示脚本和现场样例"],
        ["材料交付", "整理源码、依赖、运行说明、第三方清单与项目材料", "可复现提交包与答辩备份"],
    ], [3.0, 8.0, 6.0])

    add_heading(doc, "1 预期成果", 2)
    add_table(doc, ["成果", "形式", "评审价值"], [
        ["智能核查系统原型", "Web 工作台或命令行流程", "可现场处理样例并展示 Agent 编排"],
        ["结构化核查报告", "JSON 与可读报告", "可扫描错误、建议与证据，便于人工复核"],
        ["金标案例集", "固定草稿、可授权材料、预期结果和测试脚本", "可复跑、可验证，不依赖临场网络"],
        ["完整源码与说明", "任务编排、工具、Prompt、依赖和运行手册", "满足复现、开源依赖披露与答辩追问"],
    ], [3.4, 6.0, 7.6])
    add_heading(doc, "2 初赛阶段验收目标", 2)
    add_text(doc, "以下目标用于团队自检，不将其表述为未经验证的真实市场效果或通用模型准确率。")
    add_table(doc, ["验收项", "通过标准", "留存证据"], [
        ["主线覆盖", "稳定处理数值、单位、期间、口径、计算、估值和引用七类核心核查", "金标案例、分类结果与测试记录"],
        ["结论完整性", "每条确认错误均具备原句、来源定位、规则或公式、修改建议和状态", "结构化 JSON 与可读核查报告"],
        ["复现能力", "在固定材料、依赖与配置下重复运行，核心错误分类和证据定位保持一致", "运行日志、输入哈希与输出摘要"],
        ["审慎边界", "材料缺失、口径冲突或证据不足时不输出确认错误结论", "blocked、partial 与人工复核样例"],
        ["交付完整性", "源码、依赖、运行说明、样例、第三方清单和展示脚本齐备", "提交包目录与现场演示清单"],
    ], [3.2, 8.1, 5.7])
    add_heading(doc, "3 风险控制", 2)
    add_table(doc, ["风险", "影响", "控制措施"], [
        ["PDF 表格或扫描件解析不稳定", "事实提取不完整", "建立材料准入检查；解析失败转人工复核，不静默给出结论"],
        ["模型调用超时或输出不合规", "主张提取中断", "限制批量与重试次数；记录 partial；准备离线金标演示路径"],
        ["跨期跨口径数据误配", "产生伪错误", "事实对象强制绑定期间、单位、币种和口径；校验前拒绝混用"],
        ["第三方与数据合规", "无法提交或展示", "提交前生成第三方清单；不打包密钥、受版权限制材料或模型权重"],
    ], [4.2, 4.2, 8.6])
    add_heading(doc, "4 应用价值与后续拓展", 2)
    add_text(doc, "在研究团队、券商研究所、资管投研和企业董办等场景中，系统可以作为“研究发布前的第二道核查工序”。它不替代研究员的观点判断，而是把反复发生的事实核对、公式复算和引用追踪沉淀为可审计流程，从而缩短复核时间并降低低级错误对研究质量的影响。")
    add_text(doc, "在初赛主线稳定后，系统可扩展至公告结构化提取、财报异常信号识别、估值敏感性检查和研究报告质量评估。但任何扩展均沿用同一原则：模型负责理解与组织，金融事实、计算与结论必须能够回到来源和规则。")

    add_heading(doc, "十一 团队分工")
    add_table(doc, ["角色", "主要职责"], [
        ["Agent 与后端负责人", "任务编排 工具契约 API 运行日志"],
        ["金融规则负责人", "指标口径 错误规则 修改建议与适用边界"],
        ["文档智能负责人", "PDF 表格解析 主张提取 原文定位"],
        ["前端与演示负责人", "五页工作台 结果可视化与答辩"],
        ["测试与材料负责人", "金标案例 离线测试 计划书 第三方清单"],
    ], [4.7, 12.3])

    add_heading(doc, "十二 现场展示与答辩要点")
    add_table(doc, ["展示环节", "评委应看到的证据"], [
        ["材料登记", "输入文件类型、哈希、研究截止日与受控数据范围"],
        ["主张提取", "模型输出的逐字原句与结构化字段，而非模型直接给出财务结论"],
        ["错误核查", "数值、单位、期间、口径、计算、估值与引用的分类结果"],
        ["证据回溯", "一条错误对应的原始年报页、事实记录与程序化公式"],
        ["失败可见性", "材料缺失、模型超时或证据不足时显示 blocked 或 partial，而不是显示核验通过"],
    ], [4.0, 13.0])

    add_heading(doc, "十三 提交前验收清单")
    for item in [
        "从自然语言研报草稿中提取可核查主张，并保留逐字原句与页码。",
        "使用至少一个大语言模型作为 Agent 的任务理解和工具编排引擎。",
        "稳定覆盖数值、单位、期间、口径、计算、估值和引用等核心核查。",
        "每条确认错误均能定位研报原句和原始财务材料。",
        "每条数值类错误均能展示程序化计算过程和输入证据。",
        "证据不足时输出无法核验或待人工复核，而不是编造结论。",
        "在本地封闭或半封闭环境中完整运行；默认测试离线运行，不依赖网络、密钥、缓存或实时行情。",
        "不将合成演示数据或模型叙述表述为真实金融事实或真实模型效果；提交包不含密钥、.env 或受版权限制原始材料。",
    ]:
        add_bullet(doc, "□ " + item, compact=True)

    add_heading(doc, "十四 可复现提交包目录")
    add_text(doc, "建议以以下结构整理最终提交包，使评审人员无需依赖团队成员口头说明即可完成环境准备、样例复跑与结果核查。")
    add_table(doc, ["目录或文件", "应包含内容", "评审用途"], [
        ["README 与运行说明", "环境要求、安装步骤、启动命令、无密钥配置示例与常见失败处理", "独立复现核心流程并降低配置成本"],
        ["src 与任务编排", "Agent 编排、工具调用、金融规则、数据处理与日志模块", "核查自主开发范围与工程实现"],
        ["examples 与金标", "可授权样例材料、正确与含错草稿、预期结果和测试脚本", "离线复跑并核验误报控制"],
        ["output 示例", "结构化核查结果、可读报告、审计轨迹与证据索引", "快速理解最终业务交付物"],
        ["依赖与第三方清单", "依赖版本、模型名称与调用范围、开源项目来源及许可证", "核查环境一致性与合规性"],
    ], [3.5, 8.7, 4.8])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
